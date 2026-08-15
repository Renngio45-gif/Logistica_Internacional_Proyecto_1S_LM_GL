# -*- coding: utf-8 -*-
"""Guía de estudio y exposición del proyecto Pepa de Oro."""
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))

TINTA = RGBColor(0x24, 0x15, 0x05)
ORO = RGBColor(0xB8, 0x7A, 0x0F)
GRIS = RGBColor(0x6B, 0x57, 0x44)
VERDE = RGBColor(0x2F, 0x5D, 0x3F)
HEX_TINTA = '241505'
HEX_CLARO = 'F5F3EB'


def cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def cell_pad(cell, v=120, h=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for name, val in [('top', v), ('bottom', v), ('left', h), ('right', h)]:
        n = OxmlElement(f'w:{name}')
        n.set(qn('w:w'), str(val))
        n.set(qn('w:type'), 'dxa')
        mar.append(n)
    tcPr.append(mar)


def table_borders(table, color='CFC7B4'):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.9)
    s.bottom_margin = Inches(0.9)
    s.left_margin = Inches(0.95)
    s.right_margin = Inches(0.95)

st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
st.paragraph_format.line_spacing = 1.18
st.paragraph_format.space_after = Pt(7)


def p(text='', size=11, bold=False, italic=False, color=None,
      align=WD_ALIGN_PARAGRAPH.LEFT, space=7, indent=0, mono=False):
    par = doc.add_paragraph()
    par.alignment = align
    par.paragraph_format.space_after = Pt(space)
    if indent:
        par.paragraph_format.left_indent = Inches(indent)
    if text:
        r = par.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if mono:
            r.font.name = 'Consolas'
        if color:
            r.font.color.rgb = color
    return par


def h1(text):
    par = p(text, size=17, bold=True, color=TINTA, space=4)
    par.paragraph_format.space_before = Pt(20)
    return par


def h2(text):
    par = p(text, size=12.5, bold=True, color=ORO, space=5)
    par.paragraph_format.space_before = Pt(13)
    return par


def bullet(label, text, indent=0.22):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Inches(indent)
    par.paragraph_format.space_after = Pt(5)
    r = par.add_run('•  ')
    r.font.size = Pt(11)
    if label:
        rl = par.add_run(label)
        rl.bold = True
        rl.font.size = Pt(11)
    rt = par.add_run(text)
    rt.font.size = Pt(11)
    return par


def caja(titulo, cuerpo, fill=HEX_CLARO):
    t = doc.add_table(rows=1, cols=1)
    table_borders(t)
    c = t.rows[0].cells[0]
    c.text = ''
    cell_bg(c, fill)
    cell_pad(c, 150, 170)
    par = c.paragraphs[0]
    par.paragraph_format.space_after = Pt(3)
    r = par.add_run(titulo)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = ORO
    par2 = c.add_paragraph()
    par2.paragraph_format.space_after = Pt(0)
    r2 = par2.add_run(cuerpo)
    r2.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def guion(texto):
    """Lo que se dice en voz alta, en cursiva y con sangría."""
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Inches(0.3)
    par.paragraph_format.space_after = Pt(9)
    r = par.add_run('“' + texto + '”')
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = GRIS
    return par


# ============================== PORTADA ==============================
t = doc.add_table(rows=1, cols=1)
table_borders(t, HEX_TINTA)
c = t.rows[0].cells[0]
c.text = ''
cell_bg(c, HEX_TINTA)
cell_pad(c, 320, 260)

pp = c.paragraphs[0]
pp.paragraph_format.space_after = Pt(4)
r = pp.add_run('GUÍA DE EXPOSICIÓN')
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = ORO

pp2 = c.add_paragraph()
pp2.paragraph_format.space_after = Pt(6)
r2 = pp2.add_run('PEPA DE ORO')
r2.bold = True
r2.font.size = Pt(30)
r2.font.color.rgb = RGBColor(0xD8, 0x9A, 0x22)

pp3 = c.add_paragraph()
pp3.paragraph_format.space_after = Pt(2)
r3 = pp3.add_run('Tema 7 · Presupuesto y logística de un envío internacional')
r3.font.size = Pt(12)
r3.font.color.rgb = RGBColor(0xE8, 0xE6, 0xDC)

pp4 = c.add_paragraph()
pp4.paragraph_format.space_after = Pt(0)
r4 = pp4.add_run('Geovanny Farías Estupiñán  ·  Habilidades Lógico-Matemáticas  ·  '
                 'Msc. Adrián Vargas  ·  PUCE Sede Esmeraldas')
r4.font.size = Pt(9.5)
r4.font.color.rgb = RGBColor(0x9C, 0x7F, 0x55)

doc.add_paragraph()

p('Esta guía va de lo más simple a lo más detallado. Empieza por la página siguiente: si '
  'entiendes esa sola página, ya puedes explicar de qué se trata todo el proyecto.', size=11.5)

caja('Cómo estudiarla',
     'No memorices las fórmulas: entiende de dónde sale cada una. Si entiendes el "por qué", '
     'el "cómo" lo puedes reconstruir en el pizarrón aunque te pongan nervioso.\n\n'
     '1. Lee la página siguiente (las tres fases) hasta que se la puedas contar a alguien.\n'
     '2. Repasa el glosario y quédate con las palabras que te cuesten.\n'
     '3. Lee la Parte A dos veces.\n'
     '4. Practica la Parte B en voz alta, con reloj.')

doc.add_page_break()

# ==================== LAS TRES FASES EN LENGUAJE SIMPLE ====================
h1('Todo el proyecto en una página')

caja('La idea que ordena todo',
     'Mi proyecto responde TRES preguntas, y cada fase responde una:\n\n'
     '     FASE 1 responde CUÁNDO.\n'
     '     FASE 2 responde CÓMO.\n'
     '     FASE 3 responde CUÁNTO.\n\n'
     'Si te bloqueas en la exposición, di esa frase en voz alta. Cada palabra te lleva a una fase '
     'y recuperas el hilo.')

h2('FASE 1 · ¿CUÁNDO uso la ruta rápida?')
p('La pregunta real: tengo un lote de cacao listo. ¿Lo mando por la vía rápida y cara, o por la '
  'lenta y barata?', space=5)
p('La herramienta: lógica proposicional.', bold=True, space=5)
p('Hay dos cosas que pueden ser sí o no: ¿es urgente? (p) y ¿me pasé del presupuesto? (q). '
  'La regla dice que mando por la rápida solo si es urgente Y no me pasé de plata.', space=5)
p('D = p ∧ ¬q', size=13, bold=True, color=TINTA, mono=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=5)
p('Como son dos preguntas de sí o no, hay 4 combinaciones posibles. Solo una aprueba: cuando sí '
  'es urgente y no me pasé. Las otras tres se van por la ruta normal.', space=5)
caja('Para entenderlo con algo de todos los días',
     '¿Tomo taxi o bus? Tomo taxi solo si voy tarde Y tengo plata.\n'
     'Si voy tarde pero no tengo plata, bus. Si tengo plata pero no voy tarde, bus también.\n'
     'Esa es exactamente la regla de mi proyecto.')

h2('FASE 2 · ¿CÓMO reparto el dinero de cada envío?')
p('La pregunta real: tengo $3,500 para gastar en un envío. ¿Cuánto va a transporte y cuánto a '
  'seguro?', space=5)
p('La herramienta: sistema de dos ecuaciones.', bold=True, space=5)
p('Sé dos cosas y con eso me alcanza: que los dos juntos suman 3 500, y que el transporte cuesta '
  '4 veces el seguro. Reemplazo la segunda en la primera y todo queda en función del seguro.', space=5)
p('T + S = 3 500        T = 4S\n4S + S = 3 500  →  5S = 3 500  →  S = 700\nT = 4(700) = 2 800',
  size=12, bold=True, color=TINTA, mono=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=5)
p('Resultado: $2 800 de transporte y $700 de seguro.', bold=True, space=5)
caja('Por qué hacen falta DOS ecuaciones',
     'Con una sola no alcanza. Si solo sé que suman 3 500, podría ser 3 000 y 500, o 2 000 y 1 500, '
     'o mil combinaciones más. La segunda ecuación es la que fija cuál de todas es.\n\n'
     'Con algo de todos los días: dos personas juntan $100 y una puso el triple que la otra. '
     '¿Cuánto puso cada una? Es el mismo procedimiento.')

h2('FASE 3 · ¿CUÁNTO cacao exporto al mes?')
p('La pregunta real: ¿mando 2 contenedores? ¿5? ¿10? Mientras más mando más gano… ¿o no?', space=5)
p('La herramienta: función cuadrática (una parábola).', bold=True, space=5)
p('Acá está el truco del proyecto. Si el precio de compra fuera fijo, mientras más mandara más '
  'ganaría, sin límite. Pero no es fijo: para llenar más contenedores tengo que comprarle a fincas '
  'más lejanas y pagar más caro. Entonces vendo a precio fijo pero compro a precio que sube, y de '
  'esa pelea sale una parábola con un punto más alto.', space=5)
p('U(x) = −2 500x² + 23 550x − 20 000', size=13, bold=True, color=TINTA, mono=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, space=5)
p('Ese punto más alto cae en 4.71 contenedores. Como no existen contenedores partidos, pruebo 4 y '
  '5: gana 5, con $35 250 al mes.', space=5)
caja('Para entenderlo con algo de todos los días',
     'Vendes empanadas. Con pocas no cubres el gas ni la harina, así que pierdes. '
     'Con demasiadas se te queman o te sobran, y también pierdes. '
     'Hay un número justo en el medio donde ganas más.\n\n'
     'Esa es tu parábola: no es que "más siempre sea mejor".')

h2('Cómo se conectan las tres')
p('No son tres trabajos separados: son tres decisiones encadenadas de la misma operación.', space=5)
bullet('La Fase 3 me dice ', 'que mande 5 contenedores al mes.')
bullet('La Fase 2 me dice ', 'que cada uno lleva $2 800 de transporte y $700 de seguro.')
bullet('La Fase 1 me dice ', 'si ese envío en particular va por la ruta rápida o por la normal.',
       indent=0.22)

doc.add_page_break()

# ============================== GLOSARIO ==============================
h1('Glosario · las palabras que tienes que manejar')

p('Si el profesor usa alguna de estas palabras y no la reconoces, te bloqueas. '
  'Léelas una vez y quédate con las que te cuesten.', size=11, italic=True, color=GRIS, space=10)


def glosario(titulo, filas):
    h2(titulo)
    t = doc.add_table(rows=1, cols=2)
    table_borders(t)
    hdr = t.rows[0].cells
    for i, txt in enumerate(['Palabra', 'Qué significa y dónde aparece en tu proyecto']):
        hdr[i].text = ''
        cell_bg(hdr[i], HEX_TINTA)
        cell_pad(hdr[i], 100, 130)
        rr = hdr[i].paragraphs[0].add_run(txt)
        rr.bold = True
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = RGBColor(0xD8, 0x9A, 0x22)
    for i, (k, v) in enumerate(filas):
        row = t.add_row().cells
        fill = 'FFFFFF' if i % 2 == 0 else HEX_CLARO
        for j, txt in enumerate([k, v]):
            row[j].text = ''
            cell_bg(row[j], fill)
            cell_pad(row[j], 100, 130)
            par = row[j].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            rr = par.add_run(txt)
            rr.font.size = Pt(9.5)
            if j == 0:
                rr.bold = True
                rr.font.color.rgb = TINTA
    t.columns[0].width = Inches(1.5)
    t.columns[1].width = Inches(5.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


glosario('Palabras de lógica', [
    ('Proposición',
     'Una frase que solo puede ser verdadera o falsa, sin término medio. "El envío es urgente" '
     'sí lo es. "El peso de la carga" no, porque es un número.'),
    ('Conjunción  ∧',
     'Quiere decir "y". Para que sea verdadera se tienen que cumplir las DOS cosas. '
     'En tu regla: urgente Y con presupuesto.'),
    ('Negación  ¬',
     'Quiere decir "no". Invierte el valor: si q es verdadera, ¬q es falsa. '
     'En tu regla ¬q significa "el presupuesto NO está excedido".'),
    ('Tabla de verdad',
     'La lista de todos los casos posibles y su resultado. Con 2 proposiciones son 4 filas, '
     'porque 2² = 4. Con 3 serían 8.'),
    ('Tautología',
     'Una fórmula que sale verdadera en TODAS las filas, siempre. Por eso el Modus Ponens es '
     'confiable: nunca falla.'),
    ('Inferencia',
     'Sacar una conclusión a partir de cosas que ya sabes. Es el paso de las premisas a la '
     'conclusión.'),
    ('Modus Ponens',
     'La inferencia más simple: si sé que "A implica B" y compruebo que A pasó, entonces B pasa. '
     'Si llueve la cancha se moja; llovió; entonces se mojó.'),
])

glosario('Palabras de álgebra y funciones', [
    ('Incógnita',
     'El valor que no conozco y quiero encontrar. En tu sistema son T (transporte) y S (seguro).'),
    ('Sistema de ecuaciones',
     'Dos o más ecuaciones que se cumplen al mismo tiempo. Hacen falta tantas ecuaciones como '
     'incógnitas para tener una sola respuesta.'),
    ('Método de sustitución',
     'Despejar una incógnita en una ecuación y meterla en la otra, para quedarte con una sola '
     'incógnita. Es el que usaste.'),
    ('Comprobar',
     'Reemplazar la respuesta en las ecuaciones originales para ver que sí se cumplen. '
     'La rúbrica lo pide y vale puntos.'),
    ('Función cuadrática',
     'Una función donde la incógnita está elevada al cuadrado. Su dibujo siempre es una parábola.'),
    ('Parábola',
     'La curva en forma de U (o de U al revés). La tuya está al revés porque el coeficiente de x² '
     'es negativo.'),
    ('Coeficiente',
     'El número que multiplica a una letra. En −2 500x², el coeficiente es −2 500.'),
    ('Concavidad',
     'Hacia dónde abre la parábola. Si el coeficiente de x² es positivo abre hacia arriba y tiene '
     'mínimo; si es negativo abre hacia abajo y tiene máximo.'),
    ('Vértice',
     'El punto más alto (o más bajo) de la parábola. Se calcula con x = −b / 2a. '
     'El tuyo es el punto de máxima ganancia.'),
    ('Raíces',
     'Donde la curva cruza el eje horizontal, o sea donde el resultado es cero. '
     'En tu caso: donde no ganas ni pierdes.'),
    ('Discriminante',
     'El b² − 4ac de la fórmula general. Te dice cuántas raíces hay: si es positivo 2, '
     'si es cero 1, si es negativo ninguna.'),
    ('Dominio',
     'Todos los valores que puede tomar la x. El matemático es infinito; el tuyo, real, son '
     'contenedores enteros.'),
    ('Rango',
     'Todos los valores que puede tomar el resultado. El tuyo llega como máximo al vértice.'),
    ('Monotonía',
     'Dónde la curva sube y dónde baja. La tuya crece hasta 4.71 y decrece después.'),
])

glosario('Palabras de comercio y logística', [
    ('FOB',
     'Free On Board. El valor de la mercancía puesta a bordo del buque, sin flete ni seguro. '
     'Solo se usa para transporte por barco.'),
    ('FCA',
     'Free Carrier. El equivalente al FOB cuando el envío va por avión o por carretera. '
     'Por eso tu calculadora cambia el nombre según la vía.'),
    ('Incoterm',
     'Las reglas internacionales que definen hasta dónde llega la responsabilidad del vendedor. '
     'FOB y FCA son dos de ellas.'),
    ('Arancel',
     'El impuesto que cobra el país de destino por dejar entrar la mercancía. '
     'En la Unión Europea el cacao ecuatoriano paga 0%.'),
    ('Flete',
     'Lo que cobra la naviera o la aerolínea por mover la carga.'),
    ('Quintal',
     'La unidad tradicional del cacao en Ecuador: 100 libras, o sea 45.36 kg. '
     'Un productor piensa en quintales, no en kilos.'),
    ('Acopio',
     'Juntar el cacao de varios productores para llenar un contenedor. '
     'Mientras más juntas, más lejos tienes que ir y más caro te sale.'),
    ('Valor CIF',
     'El valor de la mercancía incluyendo ya el flete y el seguro. Se usa de base para calcular '
     'el seguro de carga.'),
])

glosario('Palabras del modelo', [
    ('Supuesto',
     'Algo que asumes porque no tienes el dato exacto. No es trampa si lo declaras. '
     'El tuyo es que el precio sube $100/t por contenedor.'),
    ('Rendimientos decrecientes',
     'Que cada unidad extra rinde menos que la anterior. Es lo que hace que tu curva baje '
     'después del vértice.'),
    ('Análisis de sensibilidad',
     'Probar qué pasaría con la respuesta si el supuesto fuera distinto. '
     'Es lo que blinda tu modelo ante la pregunta difícil.'),
    ('Punto de equilibrio',
     'Donde no ganas ni pierdes. Son tus raíces: 0.94 y 8.48 contenedores.'),
    ('Optimizar',
     'Buscar el mejor valor posible. Tú optimizas la ganancia buscando el vértice.'),
])

# ============================== PARTE A ==============================
h1('PARTE A · Lo que tienes que entender')

h2('1. El problema en una frase')

p('Un productor de cacao en Esmeraldas tiene que decidir tres cosas cada vez que exporta: '
  'si paga una ruta rápida o una barata, cómo reparte el dinero del envío, y cuánto cacao '
  'mandar. Hoy lo decide a ojo. Mi proyecto convierte esas tres decisiones en matemáticas.')

caja('La idea que amarra todo el proyecto',
     'Son tres preguntas y cada área de la materia responde una:\n\n'
     '   La LÓGICA responde CUÁNDO conviene la ruta rápida.\n'
     '   El ÁLGEBRA responde CÓMO repartir el dinero de cada envío.\n'
     '   Las FUNCIONES responden CUÁNTO cacao exportar al mes.\n\n'
     'Si te preguntan cualquier cosa y te bloqueas, vuelve a esta frase. Ordena todo.')

h2('2. La lógica, explicada fácil')

p('Una proposición es una frase que solo puede ser verdadera o falsa. No hay término medio. '
  '"El envío es urgente" es una proposición: o lo es o no lo es. En cambio "el peso de la carga" '
  'no es una proposición, porque es un número, no un sí o un no.')

bullet('p = ', 'el envío es urgente.')
bullet('q = ', 'el presupuesto está excedido.')
bullet('r = ', 'el peso total de la carga. Es un número, no una proposición: por eso no aparece '
               'en la tabla de verdad, sino que alimenta los cálculos.')

p('La regla del sistema es:', space=3)
p('D = p ∧ ¬q', size=15, bold=True, color=TINTA, mono=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=5)

p('El símbolo ∧ significa "y". El símbolo ¬ significa "no". Entonces se lee: '
  'apruebo la ruta rápida cuando el envío es urgente Y el presupuesto NO está excedido. '
  'Las dos condiciones tienen que cumplirse a la vez.')

caja('Por qué la tabla tiene exactamente 4 filas',
     'Porque hay 2 proposiciones y cada una puede tomar 2 valores (verdadero o falso). '
     'Entonces las combinaciones son 2 × 2 = 4. Si hubiera 3 proposiciones serían 2³ = 8 filas. '
     'Esa es la pregunta fácil que te puede hacer el profesor: la respuesta es 2 elevado al '
     'número de proposiciones.')

p('De las cuatro combinaciones, solo una aprueba: cuando p es verdadera y q es falsa. '
  'Tiene sentido, porque es el único caso donde de verdad hay urgencia y además alcanza el dinero.')

h2('3. El Modus Ponens, sin susto')

p('Modus Ponens es la regla de inferencia más simple que existe. Dice: si tengo una promesa '
  'del tipo "si pasa A entonces pasa B", y compruebo que A efectivamente pasó, entonces B pasa. '
  'Es el razonamiento que usas todos los días sin darte cuenta.')

caja('Ejemplo cotidiano para entenderlo',
     'Si llueve, la cancha se moja. Llovió. Entonces la cancha se mojó.\n\n'
     'En mi proyecto es igual:\n'
     'Si el envío es urgente y hay presupuesto, entonces se aprueba la ruta rápida.\n'
     'El envío es urgente y hay presupuesto.\n'
     'Entonces se aprueba la ruta rápida.')

p('Se dice que la inferencia es válida porque su forma [(A→B) ∧ A] → B es una tautología. '
  'Tautología significa que sale verdadera en todas las filas de su tabla de verdad, siempre, '
  'sin importar los valores. Por eso es imposible que las premisas sean verdaderas y la '
  'conclusión falsa.')

h2('4. El sistema de dos ecuaciones')

p('Tengo $3,500 para gastar en cada envío, repartidos entre transporte y seguro. Sé además que '
  'el transporte cuesta 4 veces lo que cuesta el seguro. Con esas dos condiciones armo el sistema:')

p('T + S = 3 500          T = 4S', size=13, bold=True, color=TINTA, mono=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, space=5)

p('T es el transporte y S es el seguro, las dos en dólares. Lo resuelvo por sustitución: '
  'como T vale 4S, reemplazo T en la primera ecuación y me queda todo en función de S.')

p('4S + S = 3 500   →   5S = 3 500   →   S = 700\n'
  'T = 4(700) = 2 800',
  size=12, mono=True, color=TINTA, align=WD_ALIGN_PARAGRAPH.CENTER, space=5)

caja('No olvides comprobar: vale puntos',
     'Reemplazo la respuesta en las DOS ecuaciones originales:\n'
     '   2 800 + 700 = 3 500  ✓\n'
     '   2 800 = 4 × 700 = 2 800  ✓\n\n'
     'Comprobar es lo que separa "resolví" de "resolví bien". La rúbrica lo pide explícitamente.')

p('Ojo con la notación: uso T y S, no x e y, para que no se confundan con la x de la función '
  'de utilidad, que cuenta contenedores. Son cosas distintas y por eso llevan letras distintas. '
  'Si el profesor te pregunta por qué, esa es la respuesta.')

h2('5. La función de utilidad: la parte que más se pregunta')

p('Aquí decido cuántos contenedores exportar al mes. Llamo x al número de contenedores.')

p('La ganancia es lo que vendo menos lo que gasto. Vender es fácil: cada contenedor lleva 25 '
  'toneladas y cada tonelada vale $5,900, así que cada contenedor me deja $147,500 de ingreso.')

p('Lo interesante es el costo de comprar el cacao. Si quiero exportar más, no me alcanza con los '
  'productores cercanos: tengo que ir a fincas más lejanas y pagar más caro para asegurar el grano. '
  'Entonces el precio no es fijo, sube con el volumen:')

p('P(x) = 4 700 + 100x     (dólares por tonelada)', size=12.5, bold=True, color=TINTA, mono=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, space=5)

caja('De aquí sale el x² — esto es lo más importante de toda la guía',
     'El costo de comprar es: número de contenedores × toneladas por contenedor × precio por tonelada.\n\n'
     '   Costo = x · 25 · (4 700 + 100x)\n'
     '   Costo = 117 500x + 2 500x²\n\n'
     'El cuadrado aparece SOLO, porque multiplico x por algo que también tiene x adentro. '
     'No lo inventé para que diera una parábola: es consecuencia de la multiplicación.\n\n'
     'Y el 2 500 tiene significado concreto: son las 25 toneladas del contenedor multiplicadas '
     'por los $100 que sube el precio. 25 × 100 = 2 500.')

p('Juntando ingreso menos todos los costos queda la función:')

p('U(x) = −2 500x² + 23 550x − 20 000', size=14, bold=True, color=TINTA, mono=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, space=6)

p('Como el número que acompaña al x² es negativo (−2 500), la parábola abre hacia abajo. '
  'Y una parábola que abre hacia abajo tiene un punto máximo. Ese máximo es exactamente lo que '
  'busco: el número de contenedores que me deja la mayor ganancia.')

h2('6. Vértice y raíces: qué significan aquí')

bullet('El vértice ', 'es el punto más alto de la parábola. Se calcula con x = −b/2a. '
                      'Aquí da 4.71 contenedores, con una ganancia de $35,460.')
bullet('Las raíces ', 'son donde la parábola cruza el eje horizontal, o sea donde la ganancia es cero: '
                      'ni gano ni pierdo. Dan 0.94 y 8.48.')

caja('El detalle que impresiona: el dominio real',
     'El vértice dice 4.71 contenedores. Pero 0.71 de un contenedor no existe. '
     'Matemáticamente la función es continua, pero en la realidad solo puedo mandar contenedores enteros.\n\n'
     'Entonces evalúo los dos enteros vecinos:\n'
     '   U(4) = $34,200\n'
     '   U(5) = $35,250  ←  gana este\n\n'
     'Por eso la respuesta final es 5 contenedores y no 4.71. Esto es lo que se llama '
     '"dominio con sentido real" y es justo lo que pedía el deber de la clase 9.')

p('Entre las raíces está la zona donde gano dinero: de 1 a 8 contenedores. Con menos de 1 no '
  'cubro los costos fijos del mes, y con más de 8 el encarecimiento del acopio se come toda la ganancia.')

doc.add_page_break()

# ============================== PARTE B ==============================
h1('PARTE B · El guion de la exposición')

caja('Antes de empezar',
     'Abre la página web en una pestaña ANTES de pasar al frente, con la calculadora ya cargada. '
     'Ten el PowerPoint en otra pestaña. Si falla el internet, las diapositivas solas te sostienen '
     'la nota: la web es un plus, no un salvavidas.')

# --- Diapositiva 1 ---
h2('Diapositiva 1 · Portada  (30 segundos)')
p('Solo la lees y te presentas. No te quedes mucho aquí.', italic=True, color=GRIS, space=5)
guion('Buenos días. Soy Geovanny Farías. Mi proyecto es el tema 7, presupuesto y logística de un '
      'envío internacional, y lo apliqué a un caso de mi provincia: la exportación de cacao fino de '
      'aroma desde Esmeraldas. Construí una herramienta web que se llama Pepa de Oro, que es como '
      'se le dice al cacao ecuatoriano.')

# --- Diapositiva 2 ---
h2('Diapositiva 2 · El problema y la lógica  (2 minutos)')
p('Señala la tabla de verdad cuando llegues a ella. Es tu evidencia del 30% de lógica.',
  italic=True, color=GRIS, space=5)
guion('El problema real es que las asociaciones cacaoteras deciden a ojo cómo exportar: pagan una '
      'ruta urgente que no necesitaban, o mueven la carga tan lento que el grano pierde calidad.\n\n'
      'Definí tres variables. p y q son proposiciones, o sea que solo pueden ser verdaderas o falsas, '
      'y por eso son las que entran en la tabla. r es numérica: no se enciende ni se apaga, sino que '
      'alimenta los cálculos del valor y del número de contenedores.\n\n'
      'La regla es D igual a p y no q: la ruta rápida se aprueba solo si el envío es urgente y el '
      'presupuesto no está excedido. Como son dos proposiciones, hay cuatro combinaciones posibles, '
      'y como ven aquí, una sola aprueba.\n\n'
      'Para validar la regla apliqué Modus Ponens: si se cumple el condicional y se cumple el '
      'antecedente, la conclusión es necesaria. Su forma es una tautología, así que si las premisas '
      'son verdaderas la conclusión no puede ser falsa.')

# --- Diapositiva 3 ---
h2('Diapositiva 3 · El modelo algebraico  (2 minutos y medio)')
p('Es la diapositiva más densa. Ve despacio y apóyate en la gráfica.',
  italic=True, color=GRIS, space=5)
guion('Con álgebra resuelvo dos cosas.\n\n'
      'La primera es cómo reparto el presupuesto de cada envío. Tengo 3 500 dólares y sé que el '
      'transporte cuesta cuatro veces el seguro. Eso me da un sistema de dos ecuaciones que resolví '
      'por sustitución: 2 800 de transporte y 700 de seguro. Lo comprobé en las dos ecuaciones.\n\n'
      'La segunda es cuántos contenedores exportar al mes. Aquí el supuesto es que comprar más cacao '
      'lo encarece, porque toca ir a fincas más lejanas: el precio sube 100 dólares por tonelada por '
      'cada contenedor adicional. Cuando multiplico los contenedores por ese precio, el x cuadrado '
      'aparece solo. No lo inventé: sale de la multiplicación, y el coeficiente 2 500 son las 25 '
      'toneladas por los 100 dólares.\n\n'
      'La parábola abre hacia abajo, entonces tiene un máximo. El vértice cae en 4.71, pero '
      'contenedores fraccionarios no existen, así que evalué los enteros vecinos: con 4 gano 34 200 '
      'y con 5 gano 35 250. Por eso la respuesta es cinco contenedores.')

# --- Diapositiva 4 ---
h2('Diapositiva 4 · La herramienta y la conclusión  (1 minuto y medio)')
p('Aquí cierras y das paso a la demostración en vivo.', italic=True, color=GRIS, space=5)
guion('Todo esto no se queda en el papel: lo programé en una página web. El usuario elige su '
      'provincia, cuánto cacao tiene, a qué país lo manda y cómo lo transporta, y la página le '
      'devuelve esta marca de embarque, que es el formato que se usa en el oficio.\n\n'
      'Ahí están también las conversiones entre el sistema internacional y el imperial. Usé el '
      'quintal, que es 100 libras o 45.36 kilos, porque es la unidad con la que un productor '
      'ecuatoriano realmente piensa: nadie dice "tengo 25 000 kilos", dicen "tengo 551 quintales".\n\n'
      'Todos los datos tienen fuente: el precio del cacao es de Anecacao, la producción de Esmeraldas '
      'del Ministerio de Agricultura, el flete del índice Drewry.\n\n'
      'La conclusión es que conviene exportar cinco contenedores al mes. Y para cerrar, se las muestro '
      'funcionando.')

# --- Demo ---
h2('Demostración en vivo  (90 segundos)')
p('La página vale 5 de los 9 puntos. Hazla al final: si algo falla, las diapositivas ya te '
  'aseguraron la nota.', italic=True, color=GRIS, space=6)

pasos_demo = [
    ('Abre la pestaña Calculadora.', 'Di: "supongamos que soy un productor de Manabí con 300 quintales".'),
    ('Elige Manabí, escribe 300, unidad quintales, destino Países Bajos.',
     'Señala cómo aparece el acarreo interno desde Portoviejo.'),
    ('Señala la marca de embarque.', 'Lee en voz alta: GYE flecha RTM, el costo y los días.'),
    ('Cambia a Aéreo.', 'Di: "miren cómo el costo se dispara pero el tránsito baja a dos días. '
                        'Ahí es donde entra la regla lógica: solo vale la pena si de verdad es urgente".'),
    ('Cambia el destino a México.', 'Muestra el arancel del 20%: "por eso la Unión Europea es mejor '
                                    'mercado, allá el cacao entra con arancel cero".'),
    ('Toca las pestañas de las fases.', 'Cierra: "y acá está toda la matemática que sostiene el cálculo".'),
]
for i, (accion, decir) in enumerate(pasos_demo, 1):
    bullet(f'Paso {i}. ', f'{accion}  {decir}')

doc.add_page_break()

# ============================== PARTE C ==============================
h1('PARTE C · Preguntas que te pueden hacer')

preguntas = [
    ('¿De dónde sale el 2 500?  (la más probable)',
     'Son las 25 toneladas que lleva un contenedor multiplicadas por los 100 dólares que sube el '
     'precio por tonelada. El 100 es una estimación mía, no un dato publicado, y así lo declaré '
     'como limitación del modelo. Por eso incluí análisis de sensibilidad: si fuera 50, el óptimo '
     'serían 9 contenedores; si fuera 150, serían 3. La estructura de la conclusión se mantiene, '
     'lo que cambia es la magnitud.'),
    ('¿Por qué el vértice da 4.71 si no existen contenedores fraccionarios?',
     'Porque el dominio matemático de la función es continuo, pero el dominio real es entero. '
     'Por eso evalué U(4) y U(5) y me quedé con el mayor, que es 5.'),
    ('¿Por qué la tabla de verdad tiene 4 filas?',
     'Porque son 2 proposiciones y cada una toma 2 valores: 2 elevado a 2 es 4. Con tres '
     'proposiciones serían 8.'),
    ('¿Por qué usas T y S en vez de x e y?',
     'Para no confundirlas con la x de la función de utilidad, que cuenta contenedores. '
     'Son magnitudes distintas, así que llevan nombres distintos.'),
    ('¿Por qué FOB en barco y FCA en avión?',
     'Porque los Incoterms 2020 reservan FOB, FAS, CFR y CIF exclusivamente para transporte '
     'marítimo y fluvial. Para aéreo y terrestre el término correcto es FCA, Free Carrier.'),
    ('¿Qué es el valor FOB?',
     'Es el valor de la mercancía puesta a bordo del buque, sin incluir el flete internacional '
     'ni el seguro. En mi lote son 25 toneladas por 5 900 dólares: 147 500 dólares.'),
    ('¿Por qué el arancel no está sumado en el costo total?',
     'Porque el arancel lo paga el importador cuando la carga llega a destino, no el exportador '
     'al despachar. Por eso lo muestro aparte, como información.'),
    ('¿Quién hizo la página web?',
     'Contesta con la verdad y demuestra que la entiendes: ofrécele explicar cualquier parte del '
     'cálculo que él quiera señalar. El enunciado decía que no era necesario programar, así que la '
     'web es trabajo adicional, no un requisito.'),
]

for q, a in preguntas:
    p(q, size=11.5, bold=True, color=TINTA, space=3)
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Inches(0.25)
    par.paragraph_format.space_after = Pt(11)
    r = par.add_run(a)
    r.font.size = Pt(10.5)
    r.font.color.rgb = GRIS

h1('Los números que debes saber de memoria')

t = doc.add_table(rows=1, cols=2)
table_borders(t)
hdr = t.rows[0].cells
for i, txt in enumerate(['Dato', 'Valor']):
    hdr[i].text = ''
    cell_bg(hdr[i], HEX_TINTA)
    cell_pad(hdr[i])
    rr = hdr[i].paragraphs[0].add_run(txt)
    rr.bold = True
    rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor(0xD8, 0x9A, 0x22)

memoria = [
    ('Regla lógica', 'D = p ∧ ¬q'),
    ('Solución del sistema', 'Transporte $2 800 · Seguro $700'),
    ('Función de utilidad', 'U(x) = −2 500x² + 23 550x − 20 000'),
    ('Vértice', '(4.71 ; $35 460)'),
    ('Respuesta final', '5 contenedores al mes → $35 250'),
    ('Precio del cacao', '$5 900 por tonelada (Anecacao, jul-2026)'),
    ('Un quintal', '100 libras = 45.36 kg'),
    ('Lote modelado', '25 toneladas = 551 quintales = $147 500 FOB'),
    ('Arancel en la Unión Europea', '0 % por el Acuerdo Multipartes'),
]
for i, (k, v) in enumerate(memoria):
    row = t.add_row().cells
    fill = 'FFFFFF' if i % 2 == 0 else HEX_CLARO
    for j, txt in enumerate([k, v]):
        row[j].text = ''
        cell_bg(row[j], fill)
        cell_pad(row[j])
        rr = row[j].paragraphs[0].add_run(txt)
        rr.font.size = Pt(10)
        if j == 0:
            rr.bold = True

doc.add_paragraph()
caja('Lo último, y es lo que más cuenta',
     'Si te bloqueas en cualquier momento, vuelve a la frase que ordena todo el proyecto: '
     'la lógica dice CUÁNDO, el sistema de ecuaciones dice CÓMO, y la función dice CUÁNTO. '
     'Con esa sola frase puedes retomar el hilo desde cualquier punto.')

salida = os.path.join(BASE, 'Guia_Exposicion_Pepa_de_Oro.docx')
doc.save(salida)
print('Creada:', salida)
