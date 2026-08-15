# -*- coding: utf-8 -*-
"""Genera el Deber de la Clase 9.1 - Analisis de Funciones (Partes A, B, C y D)."""
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
GRAF = os.path.join(BASE, 'graficas_clase9')


def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement('w:tblBorders')
    for b_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{b_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tbl_borders.append(border)
    tbl_pr.append(tbl_borders)


def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    COLOR_PRIMARY = RGBColor(27, 54, 93)
    COLOR_SECONDARY = RGBColor(70, 130, 180)
    COLOR_TEXT = RGBColor(51, 51, 51)
    COLOR_GREEN = RGBColor(15, 118, 110)
    HEX_PRIMARY = '1B365D'
    HEX_LIGHT_GREY = 'F4F6F9'

    st = doc.styles['Normal']
    st.font.name = 'Arial'
    st.font.size = Pt(11)
    st.font.color.rgb = COLOR_TEXT
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.space_after = Pt(6)

    def add_p(text="", space_after=6, line_spacing=1.15, bold=False, italic=False,
              color=None, align=WD_ALIGN_PARAGRAPH.LEFT, size=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.alignment = align
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
            if color:
                run.font.color.rgb = color
            if size:
                run.font.size = Pt(size)
        return p

    def add_heading(txt, size=13):
        p = add_p(txt, bold=True, color=COLOR_PRIMARY, space_after=10)
        p.runs[0].font.size = Pt(size)
        return p

    def add_sub(txt, size=11.5):
        p = add_p(txt, bold=True, color=COLOR_SECONDARY, space_after=6)
        p.runs[0].font.size = Pt(size)
        return p

    def add_bullet(label, text, indent=0.25, space_after=5):
        p = add_p(space_after=space_after)
        p.paragraph_format.left_indent = Inches(indent)
        r = p.add_run("•  " + label)
        r.bold = True
        p.add_run(text)
        return p

    def add_formula(txt, size=12.5, color=None, space_after=12):
        p = add_p(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=space_after)
        r = p.add_run(txt)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = color or COLOR_PRIMARY
        return p

    def add_image(fname, width=6.0, caption=None):
        path = os.path.join(GRAF, fname)
        if not os.path.exists(path):
            add_p(f"[Falta la imagen {fname}]", italic=True)
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(path, width=Inches(width))
        if caption:
            cp = add_p(caption, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                       space_after=16, size=9)
            cp.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    def styled_table(headers, rows_data, font_size=10):
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        hdr = table.rows[0].cells
        for i, title in enumerate(headers):
            hdr[i].text = ""
            p = hdr[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            r = p.add_run(title)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(font_size)
            set_cell_background(hdr[i], HEX_PRIMARY)
            set_cell_margins(hdr[i], top=150, bottom=150, left=120, right=120)
        for ri, data in enumerate(rows_data):
            row = table.add_row()
            bg = HEX_LIGHT_GREY if ri % 2 == 1 else "FFFFFF"
            for ci, text in enumerate(data):
                row.cells[ci].text = ""
                p = row.cells[ci].paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if ci == 0
                               else WD_ALIGN_PARAGRAPH.CENTER)
                r = p.add_run(text)
                r.font.size = Pt(font_size)
                set_cell_background(row.cells[ci], bg)
                set_cell_margins(row.cells[ci], top=110, bottom=110, left=120, right=120)
        return table

    # ==================== PORTADA ====================
    p = add_p("PONTIFICIA UNIVERSIDAD CATÓLICA DEL ECUADOR", bold=True,
              color=COLOR_PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, size=16)
    p.paragraph_format.space_before = Pt(36)
    add_p("SEDE ESMERALDAS", bold=True, color=COLOR_PRIMARY,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, size=14)
    add_p("PUCETEC", bold=True, color=COLOR_SECONDARY,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, size=11)
    add_p("TECNOLOGÍA EN DESARROLLO DE SOFTWARE", bold=True, color=COLOR_SECONDARY,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=48, size=11)

    p = add_p("DEBER — CLASE 9.1: ANÁLISIS DE FUNCIONES", bold=True, color=COLOR_PRIMARY,
              align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, size=16)
    p.paragraph_format.space_before = Pt(40)
    add_p("Unidad 3.2 (RdA 3, C2)\nPartes A, B, C y D — Parte D: Avance del Proyecto "
          "Integrador (Producto 2/3)\nProyecto: Presupuesto y logística de un envío "
          "internacional\n(Caso: Exportación de Cacao desde Esmeraldas)",
          bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40, size=11.5)

    p_meta = add_p(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    p_meta.paragraph_format.space_before = Pt(40)
    p_meta.paragraph_format.line_spacing = 1.3
    p_meta.add_run("Integrantes:\n").bold = True
    p_meta.add_run("Alan Cárdenas Morán\nGeovanny Farías Estupiñán\n\n")
    p_meta.add_run("Asignatura:\n").bold = True
    p_meta.add_run("Habilidades Lógico-Matemáticas\n\n")
    p_meta.add_run("Docente:\n").bold = True
    p_meta.add_run("Msc. Adrián Vargas\n\n")
    p_meta.add_run("Nivel y Período:\n").bold = True
    p_meta.add_run("Primer Nivel — Período 2026-I")

    p = add_p("Esmeraldas - Ecuador\n2026", bold=True, color=COLOR_SECONDARY,
              align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, size=11)
    p.paragraph_format.space_before = Pt(48)

    doc.add_page_break()

    # ==================== PARTE A ====================
    add_p("DEBER — ANÁLISIS DE FUNCIONES", bold=True, color=COLOR_PRIMARY,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20, size=14)

    add_heading("PARTE A — ANALIZA Y GRAFICA:  f(x) = x² − 2x − 3")

    add_sub("1) Identificación de coeficientes")
    add_p("La función es cuadrática de la forma f(x) = ax² + bx + c, con "
          "a = 1, b = −2 y c = −3. Como a = 1 > 0, la parábola abre hacia arriba, "
          "por lo que su vértice será un mínimo.", space_after=10)

    add_sub("2) Raíces (cortes con el eje x)")
    add_p("Se calcula primero el discriminante Δ = b² − 4ac:", space_after=6)
    add_formula("Δ = (−2)² − 4(1)(−3) = 4 + 12 = 16", space_after=8)
    add_p("Como Δ = 16 > 0, la parábola tiene dos raíces reales distintas. "
          "Aplicando la fórmula general:", space_after=6)
    add_formula("x = [ −b ± √Δ ] / 2a = [ 2 ± √16 ] / 2 = (2 ± 4) / 2", space_after=8)
    add_bullet("x₁ = ", "(2 − 4) / 2 = −1")
    add_bullet("x₂ = ", "(2 + 4) / 2 = 3", space_after=8)
    add_p("Comprobación por factorización: x² − 2x − 3 = (x + 1)(x − 3) = 0 ✓  "
          "Los cortes con el eje x son los puntos (−1, 0) y (3, 0).", space_after=12)

    add_sub("3) Vértice")
    add_formula("x_v = −b / 2a = −(−2) / 2(1) = 2/2 = 1", space_after=6)
    add_formula("y_v = f(1) = (1)² − 2(1) − 3 = 1 − 2 − 3 = −4", space_after=8)
    add_p("El vértice es V(1, −4) y, por abrir hacia arriba, es un MÍNIMO.", space_after=12)

    add_sub("4) Dominio y rango")
    add_bullet("Dominio: ", "todos los números reales, Dom(f) = ℝ = (−∞, +∞), porque "
                            "es un polinomio y no existe ninguna restricción para x.")
    add_bullet("Rango: ", "Rec(f) = [−4, +∞), porque el punto más bajo de la parábola "
                          "es el vértice y = −4 y desde ahí la curva sube "
                          "indefinidamente.", space_after=12)

    add_sub("5) Monotonía (dónde crece y dónde decrece)")
    add_bullet("Decrece en: ", "(−∞, 1) — antes del vértice la curva baja.")
    add_bullet("Crece en: ", "(1, +∞) — después del vértice la curva sube.", space_after=12)

    add_sub("6) Bosquejo")
    add_image('parte_A.png', 6.0,
              'Corte con el eje y en (0, −3). Cortes con el eje x en (−1, 0) y (3, 0). '
              'Vértice V(1, −4) marcado como mínimo.')

    doc.add_page_break()

    # ==================== PARTE B ====================
    add_heading("PARTE B — LA LINEAL:  f(x) = −3x + 9")

    add_sub("1) Identificación")
    add_p("Es una función lineal (afín) de la forma f(x) = mx + b, con pendiente "
          "m = −3 e intercepto b = 9. Su gráfica es una recta.", space_after=10)

    add_sub("2) Raíz (corte con el eje x)")
    add_p("Se iguala la función a cero y se despeja x:", space_after=6)
    add_formula("−3x + 9 = 0   →   −3x = −9   →   x = −9 / −3 = 3", space_after=8)
    add_p("Comprobación: f(3) = −3(3) + 9 = −9 + 9 = 0 ✓  La raíz es x = 3, "
          "es decir el punto (3, 0). Además, el corte con el eje y es "
          "f(0) = 9, o sea el punto (0, 9).", space_after=12)

    add_sub("3) Monotonía")
    add_p("Como la pendiente m = −3 es negativa (m < 0), la función es DECRECIENTE "
          "en todo su dominio: por cada unidad que avanza x, el valor de f(x) "
          "baja 3 unidades. Una recta no tiene vértice.", space_after=12)

    add_sub("4) Dominio y rango")
    add_bullet("Dominio: ", "Dom(f) = ℝ = (−∞, +∞).")
    add_bullet("Rango: ", "Rec(f) = ℝ = (−∞, +∞), porque al ser una recta no "
                          "horizontal, toma todos los valores reales.", space_after=12)

    add_sub("5) Bosquejo")
    add_image('parte_B.png', 6.0,
              'Recta decreciente que corta el eje y en (0, 9) y el eje x en (3, 0).')

    doc.add_page_break()

    # ==================== PARTE C ====================
    add_heading("PARTE C — INTERPRETA:  altura de un dron,  h(t) = −t² + 6t")

    add_p("La altura h (en metros) de un dron después de t segundos está dada por "
          "h(t) = −t² + 6t. Aquí a = −1, b = 6 y c = 0. Como a = −1 < 0, la parábola "
          "abre hacia abajo, por lo que su vértice es un MÁXIMO: existe un instante "
          "de altura máxima.", space_after=12)

    add_sub("1) ¿Cuándo vuelve al suelo? (raíces)")
    add_p("El dron está en el suelo cuando su altura es cero, es decir h(t) = 0:",
          space_after=6)
    add_formula("−t² + 6t = 0   →   t(−t + 6) = 0", space_after=8)
    add_bullet("t₁ = 0 s: ", "instante inicial, el dron despega desde el suelo.")
    add_bullet("t₂ = 6 s: ", "el dron vuelve a tocar el suelo.", space_after=8)
    add_p("Comprobación: h(6) = −(6)² + 6(6) = −36 + 36 = 0 ✓", space_after=12)

    add_sub("2) Instante de altura máxima (vértice)")
    add_formula("t_v = −b / 2a = −6 / 2(−1) = −6 / −2 = 3", space_after=6)
    add_formula("h(3) = −(3)² + 6(3) = −9 + 18 = 9", space_after=8)
    add_p("El vértice es V(3, 9): el dron alcanza su altura máxima de 9 metros a "
          "los 3 segundos de vuelo, exactamente en la mitad del recorrido "
          "(entre t = 0 y t = 6).", space_after=12)

    add_sub("3) Dominio y rango con sentido real")
    add_bullet("Dominio real: ", "t ∈ [0, 6] segundos. Aunque matemáticamente la "
                                 "parábola existe para todo t ∈ ℝ, no tiene sentido "
                                 "físico un tiempo negativo ni una altura bajo el suelo.")
    add_bullet("Rango real: ", "h ∈ [0, 9] metros, desde el suelo hasta la altura "
                               "máxima del vértice.", space_after=12)

    add_sub("4) Monotonía")
    add_bullet("Crece (sube) en: ", "(0, 3) — el dron asciende durante los primeros "
                                    "3 segundos.")
    add_bullet("Decrece (baja) en: ", "(3, 6) — el dron desciende hasta tocar el suelo.",
               space_after=12)

    add_sub("5) Bosquejo")
    add_image('parte_C.png', 6.0,
              'Trayectoria del dron: despega en t = 0 s, altura máxima de 9 m en '
              't = 3 s y aterrizaje en t = 6 s.')

    doc.add_page_break()

    # ==================== PARTE D ====================
    add_heading("PARTE D — PROYECTO: AVANCE DEL PRODUCTO 2/3")
    add_p("Proyecto: Presupuesto y logística de un envío internacional "
          "(Caso: Exportación de cacao fino de aroma desde Esmeraldas).",
          italic=True, space_after=12)

    add_sub("1) ¿Qué modela la función?")
    add_p("En las fases anteriores el proyecto modeló la decisión de ruta con lógica "
          "(D = p ∧ ¬q) y el reparto del presupuesto logístico con un sistema de "
          "ecuaciones 2×2 (transporte x = $2,800 y seguro y = $700). Ahora se modela "
          "la variable que la asociación exportadora realmente puede decidir cada mes: "
          "CUÁNTOS CONTENEDORES EXPORTAR.", space_after=8)
    add_p("Se define la función de utilidad neta mensual U(x), donde x es el número de "
          "contenedores de 25 toneladas de cacao exportados en el mes.", space_after=12)

    add_sub("2) Supuesto de modelado que genera la curva")
    add_p("Para exportar más volumen, la asociación no puede comprarle solo a los "
          "productores más cercanos: debe acopiar cacao de fincas y asociaciones cada "
          "vez más lejanas y pagar una prima creciente para asegurar el grano. Ese "
          "encarecimiento se modela como una función lineal del precio de compra por "
          "tonelada:", space_after=8)
    add_formula("P(x) = 4,700 + 100x   (dólares por tonelada)", space_after=8)
    add_p("El precio base es $4,700/t (precio en finca en Esmeraldas, aproximadamente "
          "el 80% del valor FOB) y sube $100 por tonelada por cada contenedor "
          "adicional que se quiera llenar en el mes.", space_after=8)

    p_lim = add_p(space_after=12)
    r = p_lim.add_run("Limitación del modelo: ")
    r.bold = True
    p_lim.add_run("el incremento de $100/t por contenedor es una estimación propia y no "
                  "un dato publicado. Representa los rendimientos decrecientes del "
                  "acopio: llenar más contenedores obliga a comprar en fincas cada vez "
                  "más lejanas. Por esa razón, en el punto 9 se incluye un análisis de "
                  "sensibilidad que muestra cómo cambia la recomendación si ese valor "
                  "fuera distinto.")

    p_val = add_p(space_after=8)
    r = p_val.add_run("Validación del precio base: ")
    r.bold = True
    p_val.add_run("el precio de $4,700/t supone que el productor recibe cerca del 80% "
                  "del precio de exportación. Esa proporción se contrastó con lo que "
                  "realmente se pagó en Ecuador durante 2026:")

    styled_table(
        ["Momento", "Precio internacional", "Pagado al productor", "Proporción"],
        [
            ["Enero 2026", "$4,500 / t", "$185/qq = $4,078 / t", "90.6 %"],
            ["Mayo 2026", "$4,169 / t", "$151.70/qq = $3,344 / t", "80.2 %"],
            ["Supuesto de este modelo", "$5,900 / t", "$4,700 / t", "79.7 %"],
        ],
        font_size=9.5
    )
    add_p("", space_after=8)
    add_p("El 79.7% que usa el modelo coincide con el 80.2% observado en mayo de 2026 y "
          "se ubica en el extremo conservador del rango registrado durante el año. La "
          "brecha entre ambos precios tiene además una explicación documentada: los "
          "compradores internacionales aplican un diferencial de hasta $700 por tonelada "
          "al cacao ecuatoriano y la exportación suma entre $12 y $15 por quintal en "
          "gastos. El precio base no es una cifra arbitraria, sino el resultado de una "
          "cadena de descuentos real.", space_after=14)

    add_sub("3) Derivación de la función")
    styled_table(
        ["Componente", "Cálculo", "Aporte a U(x)"],
        [
            ["Ingreso por venta (x contenedores × 25 t × $5,900/t, precio Anecacao "
             "jul-2026)", "147,500 · x", "+147,500x"],
            ["Costo de compra del cacao (x contenedores × 25 t × P(x))",
             "x · 25 · (4,700 + 100x)", "−117,500x − 2,500x²"],
            ["Costo logístico por contenedor (resultado de la Fase 2: T + S + "
             "comisiones)", "6,450 · x", "−6,450x"],
            ["Costos fijos mensuales (administración, certificación orgánica, personal)",
             "20,000", "−20,000"],
        ]
    )
    add_p("", space_after=10)
    add_p("El término cuadrático NO se asume de forma directa: aparece al multiplicar "
          "el número de contenedores por un precio que a su vez depende del número de "
          "contenedores.", space_after=8)
    add_formula("Costo de compra = x · 25 · (4,700 + 100x) = 117,500x + 2,500x²",
                size=12, space_after=8)
    add_p("Por lo tanto el coeficiente cuadrático es 25 t/contenedor × $100/t = $2,500, "
          "con unidades de dólares por contenedor al cuadrado. Agrupando el término "
          "lineal: 147,500 − 117,500 − 6,450 = $23,550 de margen por contenedor. La "
          "función queda:", space_after=8)
    add_formula("U(x) = −2,500x² + 23,550x − 20,000", size=14, space_after=12)
    add_p("Es una función cuadrática con a = −2,500 < 0, por lo que abre hacia abajo: "
          "tiene un MÁXIMO. Esto tiene sentido económico: exportar muy poco no cubre "
          "los costos fijos, pero exportar demasiado encarece el acopio y reduce la "
          "ganancia.", space_after=12)

    add_sub("4) Punto óptimo (vértice)")
    add_formula("x_v = −b / 2a = −23,550 / 2(−2,500) = 23,550 / 5,000 = 4.71", space_after=6)
    add_formula("U(4.71) = −2,500(4.71)² + 23,550(4.71) − 20,000 = $35,460.25", space_after=8)
    add_p("El vértice es V(4.71 ; $35,460.25) y es un MÁXIMO: la utilidad más alta "
          "posible del modelo.", space_after=12)

    add_sub("5) Dominio con sentido real")
    add_p("Matemáticamente la parábola existe para todo x ∈ ℝ, pero en la operación "
          "real no se puede exportar un número negativo ni fraccionario de "
          "contenedores. Por eso:", space_after=8)
    add_bullet("Dominio real: ", "x ∈ {0, 1, 2, 3, …} contenedores completos "
                                 "(números enteros no negativos).")
    add_bullet("Zona con ganancia: ", "resolviendo U(x) = 0 con la fórmula general se "
                                      "obtienen las raíces x ≈ 0.94 y x ≈ 8.48; entre "
                                      "ellas la utilidad es positiva. En enteros: de 1 "
                                      "a 8 contenedores el mes deja ganancia.")
    add_bullet("Restricción de abastecimiento: ", "Esmeraldas produce 58,965 t al año "
                                                  "(≈ 4,914 t al mes), muy por encima "
                                                  "de las 125 t que exigen 5 "
                                                  "contenedores, así que la materia "
                                                  "prima no limita el óptimo.",
               space_after=10)
    add_p("Como el vértice x_v = 4.71 no es un número entero, se evalúan los dos "
          "enteros vecinos para hallar el óptimo real:", space_after=8)

    styled_table(
        ["x (contenedores)", "0", "1", "2", "3", "4", "5", "6", "7", "8", "9"],
        [["U(x) en USD", "−20,000", "1,050", "17,100", "28,150", "34,200",
          "35,250", "31,300", "22,350", "8,400", "−10,550"]],
        font_size=8.5
    )
    add_p("", space_after=10)
    add_p("U(4) = $34,200 y U(5) = $35,250, por lo que el óptimo real es x = 5 "
          "contenedores, con una utilidad de $35,250 al mes.", space_after=12)

    add_sub("6) Monotonía (dónde crece y dónde decrece)")
    add_bullet("Crece en: ", "(0 ; 4.71) — cada contenedor adicional aumenta la "
                             "utilidad, porque el margen todavía supera al "
                             "encarecimiento del acopio.")
    add_bullet("Decrece en: ", "(4.71 ; +∞) — a partir del sexto contenedor el "
                               "encarecimiento del acopio crece más rápido que el "
                               "ingreso y la utilidad cae.", space_after=12)

    add_sub("7) Rango")
    add_bullet("Rango matemático: ", "U ∈ (−∞ ; 35,460.25], acotado por arriba por "
                                     "el vértice.")
    add_bullet("Rango en el dominio real (x de 1 a 8): ", "de $1,050 a $35,250.",
               space_after=12)

    add_sub("8) Bosquejo")
    add_image('parte_D.png', 6.3,
              'U(x) = −2500x² + 23550x − 20000. Raíces (equilibrio) en x ≈ 0.94 y '
              'x ≈ 8.48; vértice V(4.71 ; $35,460); óptimo real en x = 5 contenedores.')

    add_sub("9) Análisis de sensibilidad: ¿y si el supuesto fuera otro?")
    add_p("Como el incremento de $100/t es una estimación propia del grupo, se "
          "comprueba qué ocurre con la recomendación si ese valor cambia. Llamando m "
          "al incremento del precio por tonelada por cada contenedor adicional, el "
          "coeficiente cuadrático es 25m y el vértice se ubica en "
          "x_v = 23,550 / (50m) = 471/m:", space_after=10)

    styled_table(
        ["m (USD/t por contenedor)", "Función U(x)", "Vértice", "Óptimo entero",
         "Utilidad"],
        [
            ["$50 (acopio fácil)", "−1,250x² + 23,550x − 20,000", "9.42",
             "9 contenedores", "$90,700"],
            ["$100 (caso base)", "−2,500x² + 23,550x − 20,000", "4.71",
             "5 contenedores", "$35,250"],
            ["$150 (acopio difícil)", "−3,750x² + 23,550x − 20,000", "3.14",
             "3 contenedores", "$16,900"],
            ["$200 (escasez de grano)", "−5,000x² + 23,550x − 20,000", "2.35",
             "2 contenedores", "$7,100"],
        ],
        font_size=9
    )
    add_p("", space_after=10)
    p_sens = add_p(space_after=12)
    r = p_sens.add_run("Qué nos dice la sensibilidad: ")
    r.bold = True
    p_sens.add_run("la estructura de la conclusión es robusta, porque en los cuatro "
                   "escenarios existe un número óptimo de contenedores y exportar de "
                   "más siempre reduce la utilidad. Lo que sí cambia es la magnitud: "
                   "mientras más caro sea acopiar el cacao, menos contenedores conviene "
                   "mover. Por eso la recomendación de 5 contenedores es válida bajo el "
                   "supuesto de $100/t, y la asociación debería medir ese valor real en "
                   "campo antes de comprometer volúmenes mayores.")

    add_sub("10) Interpretación técnica del óptimo")
    add_p("El análisis de la función indica que la asociación exportadora debe enviar "
          "5 contenedores de cacao al mes (125 toneladas) para obtener la utilidad "
          "máxima de $35,250. Con menos de 1 contenedor o más de 8 la operación entra "
          "en pérdida, porque en el primer caso no se cubren los costos fijos y en el "
          "segundo el encarecimiento del acopio se come el margen.", space_after=8)
    add_p("Este resultado se conecta con las fases anteriores del proyecto: cada uno de "
          "esos 5 contenedores se despacha con el presupuesto logístico repartido según "
          "el sistema 2×2 de la Fase 2 (transporte T = $2,800 y seguro S = $700), y "
          "para cada envío se aplica la regla lógica D = p ∧ ¬q de la Fase 1 para "
          "decidir si corresponde la ruta rápida directa. Así, el análisis de funciones "
          "aporta el CUÁNTO exportar, el sistema de ecuaciones aporta el CÓMO repartir "
          "el presupuesto y la lógica aporta el CUÁNDO usar la ruta urgente.",
          space_after=12)

    add_sub("Fuentes de los datos reales")
    add_p("Precio del cacao ($5,900/t, julio 2026): Anecacao / Ecuavisa. Producción de "
          "Esmeraldas (58,965 t): Ministerio de Agricultura y Ganadería. Flete de "
          "contenedor de 40 pies: Drewry World Container Index 2026. Seguro de carga "
          "(0.3%–1.5% del CIF): Logintec / TransNatur. Arancel del cacao en la Unión "
          "Europea (0%): Acuerdo Comercial Multipartes Ecuador–UE.", size=9.5)

    out = os.path.join(BASE, 'Deber_Clase_9_Analisis_de_Funciones.docx')
    doc.save(out)
    print('Documento creado:', out)


if __name__ == '__main__':
    main()
