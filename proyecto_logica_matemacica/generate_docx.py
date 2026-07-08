import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    """Sets the background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Sets inner margins (padding) of a table cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Applies borders to the entire table."""
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement('w:tblBorders')
    for b_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{b_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tbl_borders.append(border)
    tbl_pr.append(tbl_borders)

def main():
    doc = Document()

    # Set page margins to standard 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Colors
    COLOR_PRIMARY = RGBColor(27, 54, 93)      # Deep Navy/Slate Blue (#1B365D)
    COLOR_SECONDARY = RGBColor(70, 130, 180)  # Muted Blue (#4682B4)
    COLOR_TEXT = RGBColor(51, 51, 51)         # Charcoal (#333333)
    HEX_PRIMARY = '1B365D'
    HEX_LIGHT_GREY = 'F4F6F9'

    # Configure Normal style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # Helper function for adding paragraphs with spacing
    def add_p(text="", style='Normal', space_after=6, line_spacing=1.15, bold=False, italic=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.alignment = align
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
            if color:
                run.font.color.rgb = color
        return p

    # --- PORTADA (COVER PAGE) ---
    p_uni = add_p("PONTIFICIA UNIVERSIDAD CATÓLICA DEL ECUADOR", bold=True, color=COLOR_PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    p_uni.runs[0].font.size = Pt(16)
    p_uni.paragraph_format.space_before = Pt(36)

    p_sede = add_p("SEDE ESMERALDAS", bold=True, color=COLOR_PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    p_sede.runs[0].font.size = Pt(14)

    p_pucetec = add_p("PUCETEC", bold=True, color=COLOR_SECONDARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    p_pucetec.runs[0].font.size = Pt(11)
    p_tech = add_p("TECNOLOGÍA EN DESARROLLO DE SOFTWARE", bold=True, color=COLOR_SECONDARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=48)
    p_tech.runs[0].font.size = Pt(11)

    p_title_cov = add_p("PROYECTO INTEGRADOR: FASE 1 (DEFINICIÓN)", bold=True, color=COLOR_PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    p_title_cov.runs[0].font.size = Pt(16)
    p_title_cov.paragraph_format.space_before = Pt(48)

    p_topic_cov = add_p("TEMA:\nPresupuesto y logística de un envío internacional\n(Caso: Exportación de Cacao desde Esmeraldas)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=48)
    p_topic_cov.runs[0].font.size = Pt(12)

    # Metadata block
    p_meta = add_p(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    p_meta.paragraph_format.space_before = Pt(48)
    p_meta.paragraph_format.line_spacing = 1.3
    
    r_int = p_meta.add_run("Integrantes:\n")
    r_int.bold = True
    p_meta.add_run("Alan Cárdenas Morán\nGeovanny Farías Estupiñán\n\n")
    
    r_asig = p_meta.add_run("Asignatura:\n")
    r_asig.bold = True
    p_meta.add_run("Habilidades Lógico-Matemáticas\n\n")

    r_doc = p_meta.add_run("Docente:\n")
    r_doc.bold = True
    p_meta.add_run("Msc. Adrián Vargas\n\n")

    r_per = p_meta.add_run("Nivel y Período:\n")
    r_per.bold = True
    p_meta.add_run("Primer Nivel — Período 2026-1")

    p_date = add_p("Esmeraldas - Ecuador\n2026", bold=True, color=COLOR_SECONDARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    p_date.runs[0].font.size = Pt(11)
    p_date.paragraph_format.space_before = Pt(60)

    # Page break to start the report
    doc.add_page_break()

    # --- REPORT TITLE (Page 2) ---
    p_title = add_p("REPORTE DE PROYECTO INTEGRADOR: FASE 1 (DEFINICIÓN)", bold=True, color=COLOR_PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
    p_title.runs[0].font.size = Pt(14)

    # --- SECTION 1: DATOS GENERALES ---
    p_h1 = add_p("1. DATOS GENERALES", bold=True, color=COLOR_PRIMARY, space_after=10)
    p_h1.runs[0].font.size = Pt(12.5)

    p_item1 = add_p(space_after=4)
    p_item1.paragraph_format.left_indent = Inches(0.25)
    r1 = p_item1.add_run("•  Integrantes: ")
    r1.bold = True
    p_item1.add_run("Alan Cárdenas Morán, Geovanny Farías Estupiñán")

    p_item2 = add_p(space_after=20)
    p_item2.paragraph_format.left_indent = Inches(0.25)
    r2 = p_item2.add_run("•  Idea de proyecto elegida: ")
    r2.bold = True
    p_item2.add_run("7. Presupuesto y logística de un envío internacional")

    # --- SECTION 2: DESCRIPCIÓN DEL SISTEMA ---
    p_h2 = add_p("2. DESCRIPCIÓN DEL SISTEMA Y PROBLEMA REAL QUE RESUELVE", bold=True, color=COLOR_PRIMARY, space_after=10)
    p_h2.runs[0].font.size = Pt(12.5)

    p_desc1 = add_p("El sistema modela y optimiza la toma de decisiones logísticas para la exportación de productos agrícolas locales (como el cacao fino de aroma) desde la provincia de Esmeraldas hacia mercados internacionales (como Europa o Estados Unidos).", space_after=8)
    
    p_desc2 = add_p(space_after=20)
    r_prob = p_desc2.add_run("Problema real: ")
    r_prob.bold = True
    p_desc2.add_run("Las asociaciones de productores en Esmeraldas enfrentan pérdidas económicas debido a la falta de un criterio automatizado y claro para elegir rutas de transporte. Frecuentemente se toman decisiones apresuradas sobre envíos urgentes que terminan excediendo el presupuesto asignado, o se seleccionan rutas lentas para carga perecedera. Este sistema resuelve la problemática delimitando de manera matemática cuándo se aprueba una ruta de distribución directa y urgente, balanceando la prioridad del tiempo de entrega con la viabilidad financiera de la exportación.")

    # --- SECTION 3: VARIABLES PRINCIPALES ---
    p_h3 = add_p("3. VARIABLES PRINCIPALES DEL SISTEMA", bold=True, color=COLOR_PRIMARY, space_after=10)
    p_h3.runs[0].font.size = Pt(12.5)

    add_p("El sistema evalúa el estado del envío mediante tres variables concretas adaptadas a la operación de exportación:", space_after=8)

    v1 = add_p()
    v1.paragraph_format.left_indent = Inches(0.25)
    v1.add_run("1.  Estado de Urgencia de la Carga (").bold = True
    v1.add_run("p").italic = True
    v1.runs[-1].bold = True
    v1.add_run("): ").bold = True
    v1.add_run("Determina si el producto corre riesgo de caducidad o si existe un contrato con penalización por retraso en el puerto de destino.")

    v2 = add_p()
    v2.paragraph_format.left_indent = Inches(0.25)
    v2.add_run("2.  Límite Presupuestario (").bold = True
    v2.add_run("q").italic = True
    v2.runs[-1].bold = True
    v2.add_run("): ").bold = True
    v2.add_run("Monitorea si los costos calculados de flete, aduanas e imprevistos superan el techo financiero asignado para el lote de envío.")

    v3 = add_p(space_after=20)
    v3.paragraph_format.left_indent = Inches(0.25)
    v3.add_run("3.  Peso Total del Envío (").bold = True
    v3.add_run("r").italic = True
    v3.runs[-1].bold = True
    v3.add_run("): ").bold = True
    v3.add_run("Registra la masa neta de la carga contratada, la cual afecta directamente las tarifas aeroportuarias o marítimas y el consumo estimado de combustible.")

    # --- SECTION 4: PROPOSICIONES LÓGICAS ---
    p_h4 = add_p("4. PROPOSICIONES LÓGICAS Y REGLA DE DECISIÓN", bold=True, color=COLOR_PRIMARY, space_after=10)
    p_h4.runs[0].font.size = Pt(12.5)

    add_p("Para automatizar la aprobación de la ruta directa, se definen las siguientes proposiciones atómicas:", space_after=8)

    prop1 = add_p()
    prop1.paragraph_format.left_indent = Inches(0.25)
    prop1.add_run("•  ").bold = True
    prop1.add_run("p").italic = True
    prop1.runs[-1].bold = True
    prop1.add_run(": El envío es calificado como urgente.")

    prop2 = add_p(space_after=12)
    prop2.paragraph_format.left_indent = Inches(0.25)
    prop2.add_run("•  ").bold = True
    prop2.add_run("q").italic = True
    prop2.runs[-1].bold = True
    prop2.add_run(": El presupuesto del envío está excedido.")

    dec_rule = add_p()
    r_dec = dec_rule.add_run("Regla de Decisión del Sistema (")
    r_dec.bold = True
    r_dec_sym = dec_rule.add_run("D")
    r_dec_sym.bold = True
    r_dec_sym.italic = True
    r_dec_end = dec_rule.add_run("): ")
    r_dec_end.bold = True
    dec_rule.add_run("La ruta rápida y directa se aprueba (")
    r_d_var = dec_rule.add_run("D")
    r_d_var.italic = True
    dec_rule.add_run(") si y solo si el envío es urgente y el presupuesto ")
    dec_rule.add_run("no").bold = True
    dec_rule.add_run(" está excedido.")

    # Mathematical Formula
    formula = add_p(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
    r_eq = formula.add_run("D = p ∧ ¬q")
    r_eq.bold = True
    r_eq.font.size = Pt(12.5)
    r_eq.font.color.rgb = COLOR_PRIMARY

    # Truth Table Title
    p_table_title = add_p("Tabla de Verdad de la Regla de Decisión", bold=True, color=COLOR_SECONDARY, space_after=8)
    p_table_title.runs[0].font.size = Pt(11.5)

    add_p("A continuación se detallan las 4 combinaciones lógicas posibles para el flujo del programa:", space_after=10)

    # Truth Table Creation
    # Columns: p, q, ¬q, D = p ∧ ¬q (Ruta Aprobada)
    headers = [
        "p (Urgente)",
        "q (Presupuesto Excedido)",
        "¬q",
        "D = p ∧ ¬q (Ruta Aprobada)"
    ]

    rows_data = [
        ["V", "V", "F", "F (Rechazado por costo elevado)"],
        ["V", "F", "V", "V (Aprobado: se requiere prioridad y hay fondos)"],
        ["F", "V", "F", "F (Rechazado: no es urgente y supera el costo)"],
        ["F", "F", "V", "F (Rechazado: no amerita ruta rápida, se usa ruta estándar)"]
    ]

    table = doc.add_table(rows=1, cols=4)
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        run = p.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)
        set_cell_background(hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(hdr_cells[i], top=160, bottom=160, left=120, right=120)

    # Format Data Rows
    for row_idx, data in enumerate(rows_data):
        row = table.add_row()
        cells = row.cells
        bg_color = HEX_LIGHT_GREY if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cells[col_idx].text = ""
            p = cells[col_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            
            # Align first 3 columns centered, and 4th column left aligned
            if col_idx < 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                if text in ["V", "F"]:
                    run.bold = True
                    if text == "V":
                        run.font.color.rgb = RGBColor(34, 139, 34)  # Forest Green
                    else:
                        run.font.color.rgb = RGBColor(178, 34, 34)  # Firebrick Red
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                # Split the text into the bold value and the description
                val = text[0] # V or F
                desc = text[1:] # e.g. " (Rechazado...)"
                run_val = p.add_run(val)
                run_val.bold = True
                if val == "V":
                    run_val.font.color.rgb = RGBColor(34, 139, 34)
                else:
                    run_val.font.color.rgb = RGBColor(178, 34, 34)
                p.add_run(desc)

            set_cell_background(cells[col_idx], bg_color)
            set_cell_margins(cells[col_idx], top=120, bottom=120, left=120, right=120)

    # Add space after table
    p_spacer = add_p(space_after=20)

    # --- SECTION 5: MAGNITUDES Y UNIDADES ---
    p_h5 = add_p("5. MAGNITUDES Y UNIDADES DE MEDIDA", bold=True, color=COLOR_PRIMARY, space_after=10)
    p_h5.runs[0].font.size = Pt(12.5)

    add_p("Para garantizar la compatibilidad del software con clientes norteamericanos y europeos, el sistema procesará y convertirá de forma automática las siguientes magnitudes en ambos sistemas de medida:", space_after=12)

    # Sub-items formatting
    def add_magnitude_item(title, si_val, imp_val):
        p_title = add_p(space_after=4)
        p_title.paragraph_format.left_indent = Inches(0.25)
        r_bullet = p_title.add_run("•  ")
        r_bullet.bold = True
        r_t = p_title.add_run(title)
        r_t.bold = True
        
        p_si = add_p(space_after=2)
        p_si.paragraph_format.left_indent = Inches(0.50)
        p_si.add_run("-  Sistema Internacional (SI): ").italic = True
        p_si.add_run(si_val)
        
        p_imp = add_p(space_after=8)
        p_imp.paragraph_format.left_indent = Inches(0.50)
        p_imp.add_run("-  Sistema Imperial: ").italic = True
        p_imp.add_run(imp_val)

    add_magnitude_item("Distancia de la Ruta Internacional:", "Kilómetros (km)", "Millas (mi)")
    add_magnitude_item("Volumen de Combustible Requerido:", "Litros (L)", "Galones (gal)")
    add_magnitude_item("Masa / Peso de la Carga de Cacao:", "Kilogramos (kg)", "Libras (lb)")

    # Financial indicator (bullet format)
    p_fin = add_p(space_after=12)
    p_fin.paragraph_format.left_indent = Inches(0.25)
    r_bullet = p_fin.add_run("•  ")
    r_bullet.bold = True
    r_t = p_fin.add_run("Indicadores Financieros:")
    r_t.bold = True

    p_fin_desc = add_p(space_after=12)
    p_fin_desc.paragraph_format.left_indent = Inches(0.50)
    p_fin_desc.add_run("-  Porcentaje (%) de impuestos aduaneros, aranceles de importación y comisiones logísticas aplicadas al valor FOB de la mercancía.")

    # Save document
    filename = "Reporte_Fase_1_Diseno_Completo.docx"
    doc.save(filename)
    print(f"Document successfully created: {filename}")

if __name__ == "__main__":
    main()
